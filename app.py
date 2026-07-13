"""
Voucher Filtration Automation — Streamlit Web App
==================================================
Run: streamlit run app.py

Tab 1 — Run Filtration:
  Appends ALU_NO / Live / Status / RRP / SRP / RRP check / % / Exclusions
  plus one YES/blank column per mechanic directly onto the original Lazada
  SellerPriceTemplate sheet (all original sheets, header rows, and formatting
  preserved exactly as exported).

Tab 2 — Clear Output Columns:
  Strips those appended columns from an already-processed file, returning it
  to the clean original Lazada template format.
"""

import streamlit as st
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
import io
import warnings
from datetime import date, datetime

warnings.filterwarnings('ignore')

NA = '#N/A'
TODAY = date.today()

# ── zeCOM column auto-detection ────────────────────────────────
# Candidate header text (lowercase) for each field we need to locate.
# Add more variants here if your tracker's column names ever change.
ZECOM_FIELD_CANDIDATES = {
    'style':         ['style#', 'style #', 'style', 'alu_no', 'alu no', 'aluno'],
    'launch_date':   ['launch date', 'launch_date', 'live date'],
    'status_laz':    ['status lazada', 'status_lazada', 'lazada status'],
    'price':         ['price (rrp)', 'rrp'],
    'special_price': ['special price (srp)', 'special price', 'srp'],
    'disc_pct':      ['disc %', 'disc%', 'discount %', 'disc pct'],
    'exclusion':     ['exclusion'],
}


def auto_detect_zecom_columns(zecom_file, header_scan_rows=10):
    """Returns {'header_row': int, 'cols': {field: idx}, 'score': int,
    'total_fields': int} on a confident match, else None. Never raises —
    falls back silently so the UI can show a manual-entry warning instead.

    This uses targeted, anchor-based rules rather than a generic per-row
    text scan, because the real zeCOM tracker has a two-row header (a
    merged group label like "EXCLUSION" one row above a blank sub-header
    cell), and several ambiguous repeated column names (multiple "MY RRP"
    / "Launch Date" columns for different marketplaces):
      - Style#: first column containing "style"
      - Launch Date: first column mentioning both "launch" and "lazada",
        excluding any explicitly marked "(ignore)"
      - Status Lazada: a column whose header is *exactly* "Lazada" (not
        merely containing it — avoids matching group labels like
        "Lazada, Zalora & Tiktok Only")
      - RRP / Special Price / Exclusion: anchored as fixed offsets around
        the first "DISC %" column, since this template always lays out
        RRP, SRP, Disc%, Exclusion as one contiguous 4-column block.
    """
    try:
        zecom_file.seek(0)
        wb = openpyxl.load_workbook(zecom_file, read_only=True, data_only=True)
        if 'MY' not in wb.sheetnames:
            wb.close()
            return None
        ws = wb['MY']

        # Header row = the row (within the first N) with the most non-blank
        # text cells — the row carrying per-column labels like "Style#".
        best_row, best_vals, best_count = None, None, 0
        for r in range(1, header_scan_rows + 1):
            try:
                row_vals = [c.value for c in ws[r]]
            except IndexError:
                break
            text_count = sum(1 for v in row_vals if isinstance(v, str) and v.strip())
            if text_count > best_count:
                best_count, best_row, best_vals = text_count, r, row_vals

        if best_row is None:
            wb.close()
            return None

        def norm(v):
            return str(v).strip().lower() if v is not None else ""

        row = best_vals

        style_idx = next(
            (i for i, v in enumerate(row) if any(c in norm(v) for c in ZECOM_FIELD_CANDIDATES['style'])),
            None
        )
        launch_idx = next(
            (i for i, v in enumerate(row) if 'launch' in norm(v) and 'lazada' in norm(v) and 'ignore' not in norm(v)),
            None
        )
        if launch_idx is None:  # fallback if no lazada-specific launch column exists
            launch_idx = next(
                (i for i, v in enumerate(row)
                 if any(c in norm(v) for c in ZECOM_FIELD_CANDIDATES['launch_date']) and 'ignore' not in norm(v)),
                None
            )
        status_idx = next(
            (i for i, v in enumerate(row) if norm(v) in ('lazada', 'status lazada', 'status_lazada')),
            None
        )
        disc_idx = next(
            (i for i, v in enumerate(row) if any(c in norm(v) for c in ZECOM_FIELD_CANDIDATES['disc_pct'])),
            None
        )

        price_idx = special_idx = exclusion_idx = None
        if disc_idx is not None and disc_idx >= 2:
            price_idx = disc_idx - 2
            special_idx = disc_idx - 1
            exclusion_idx = disc_idx + 1

        cols = {}
        for field, idx in [
            ('style', style_idx), ('launch_date', launch_idx), ('status_laz', status_idx),
            ('price', price_idx), ('special_price', special_idx),
            ('disc_pct', disc_idx), ('exclusion', exclusion_idx),
        ]:
            if idx is not None:
                cols[field] = idx

        wb.close()
        zecom_file.seek(0)

        # Require the anchor (disc_pct) and style to be found, plus at
        # least 5 of 7 fields overall, before trusting the result.
        if len(cols) < 5 or 'style' not in cols or 'disc_pct' not in cols:
            return None

        return {
            'header_row': best_row,
            'cols': cols,
            'score': len(cols),
            'total_fields': len(ZECOM_FIELD_CANDIDATES),
        }
    except Exception:
        return None


def extract_unique_exclusions(zecom_file, excl_idx, start_row=5):
    """Returns a sorted list of unique, non-blank Exclusion values found in
    the zeCOM 'MY' sheet, using the given 0-based column index. Never
    raises — returns [] on any failure so the UI can fall back to manual
    text entry."""
    try:
        zecom_file.seek(0)
        wb = openpyxl.load_workbook(zecom_file, read_only=True, data_only=True)
        if 'MY' not in wb.sheetnames:
            wb.close()
            return []
        ws = wb['MY']
        vals = set()
        for row in ws.iter_rows(min_row=start_row, values_only=True):
            if excl_idx < len(row):
                v = row[excl_idx]
                if v not in (None, ''):
                    vals.add(str(v).strip())
        wb.close()
        zecom_file.seek(0)
        return sorted(vals)
    except Exception:
        return []


# ── Page config ───────────────────────────────────────────────
st.set_page_config(
    page_title="Voucher Filtration",
    page_icon="🏷️",
    layout="wide"
)

st.markdown("""
<style>
    .main { background-color: #0f1117; }
    .block-container { padding-top: 2rem; }
    .stButton > button {
        background-color: #6c63ff;
        color: white;
        font-weight: 700;
        border: none;
        padding: 0.6rem 2rem;
        border-radius: 6px;
        font-size: 16px;
        width: 100%;
    }
    .stButton > button:hover { background-color: #8076ff; }
</style>
""", unsafe_allow_html=True)

st.title("🏷️ Voucher Filtration Automation")
st.caption(f"Lazada Campaign SKU Filtering  ·  Run date: {TODAY}")
st.divider()

# ── Tabs ──────────────────────────────────────────────────────
tab_run, tab_clear = st.tabs(["▶  Run Filtration", "🗑️  Clear Output Columns"])

# ============================================================
# SIDEBAR — global (visible in both tabs)
# ============================================================
with st.sidebar:
    st.header("⚙️ Settings")

    st.subheader("Eligibility Rules")
    st.caption(
        "A mechanic is only marked YES if Status = YES, the Live (launch) "
        "date isn't in the future, and the SKU isn't on the AM Exclusion list."
    )
    apply_threshold = st.checkbox("Also require RRP & SRP ≥ a minimum price", value=False)
    min_price = st.number_input(
        "Minimum price (applies to both RRP and SRP)", value=39, min_value=0, step=1, disabled=not apply_threshold
    )

    st.divider()
    st.subheader("zeCOM Column Positions")

    # The zecom file uploader lives further down in Tab 1, but its widget
    # state (key='zecom') persists in session_state across reruns, so we
    # can read it here even though the sidebar renders first.
    _zecom_file = st.session_state.get('zecom')
    _fallback_defaults = {
        'style': 2, 'launch_date': 21, 'status_laz': 23,
        'price': 48, 'special_price': 49, 'disc_pct': 50, 'exclusion': 51,
    }

    _auto_detected = None
    if _zecom_file is not None:
        _cache_key = f"_zecom_detect_{_zecom_file.name}_{_zecom_file.size}"
        if _cache_key not in st.session_state:
            st.session_state[_cache_key] = auto_detect_zecom_columns(_zecom_file)
        _auto_detected = st.session_state[_cache_key]

    auto_mode = st.checkbox(
        "Auto-detect from zeCOM file header",
        value=True,
        disabled=_zecom_file is None,
        help="Reads the header row of the uploaded zeCOM Tracking file and "
             "locates each column by name. Upload the zeCOM file in Tab 1 first."
    )

    if _zecom_file is None:
        st.caption("📄 Upload the zeCOM Tracking file below to enable auto-detection. Manual values used until then.")
    elif _auto_detected is None:
        st.warning(
            "⚠️ Couldn't confidently auto-detect columns from this zeCOM file — "
            "falling back to manual values below. Check that the sheet is named "
            "'MY' and headers are recognizable, or adjust the numbers yourself."
        )
    elif auto_mode:
        st.success(
            f"✓ Auto-detected {_auto_detected['score']}/{_auto_detected['total_fields']} "
            f"columns from header row {_auto_detected['header_row']}"
        )

    _defaults = dict(_fallback_defaults)
    if auto_mode and _auto_detected:
        _defaults.update(_auto_detected['cols'])

    # Keying on the file identity means a newly uploaded file always shows
    # its own freshly-detected defaults, while manual edits within the same
    # file session are preserved across reruns.
    _file_key = f"{_zecom_file.name}_{_zecom_file.size}" if _zecom_file is not None else "none"
    _locked = auto_mode and _auto_detected is not None
    st.caption("0-based index" + (" — locked to auto-detected values (uncheck above to edit)" if _locked else " — only change if the tracking file's layout changes"))

    col_style  = st.number_input("Style# (ALU_NO)", value=_defaults['style'], min_value=0, disabled=_locked, key=f"col_style_{_file_key}")
    col_launch = st.number_input("Launch Date", value=_defaults['launch_date'], min_value=0, disabled=_locked, key=f"col_launch_{_file_key}")
    col_laz    = st.number_input("Status Lazada", value=_defaults['status_laz'], min_value=0, disabled=_locked, key=f"col_laz_{_file_key}")
    col_price  = st.number_input("Price (RRP)", value=_defaults['price'], min_value=0, disabled=_locked, key=f"col_price_{_file_key}")
    col_sp     = st.number_input("Special Price (SRP)", value=_defaults['special_price'], min_value=0, disabled=_locked, key=f"col_sp_{_file_key}")
    col_disc   = st.number_input("Disc %", value=_defaults['disc_pct'], min_value=0, disabled=_locked, key=f"col_disc_{_file_key}")
    col_excl   = st.number_input("Exclusion", value=_defaults['exclusion'], min_value=0, disabled=_locked, key=f"col_excl_{_file_key}")

    st.divider()
    st.subheader("Price/Stock File Structure")
    data_start_row = st.number_input(
        "Data starts at row", value=5, min_value=2,
        help="Lazada's export keeps rows 2-4 as Mandatory/Optional/description rows; real data usually starts at row 5."
    )

ZECOM_COLS = {
    'style': int(col_style), 'launch_date': int(col_launch), 'status_laz': int(col_laz),
    'price': int(col_price), 'special_price': int(col_sp),
    'disc_pct': int(col_disc), 'exclusion': int(col_excl),
}

# Fetch the actual unique Exclusion values seen in the uploaded zeCOM file,
# so mechanics can be built by picking from real data instead of guessing
# wording. Falls back to an empty list (→ manual text entry) if no file
# is uploaded yet or extraction fails.
UNIQUE_EXCLUSIONS = []
if _zecom_file is not None:
    _excl_start_row = (_auto_detected['header_row'] + 1) if (auto_mode and _auto_detected) else 5
    _excl_cache_key = f"_zecom_excl_{_zecom_file.name}_{_zecom_file.size}_{ZECOM_COLS['exclusion']}_{_excl_start_row}"
    if _excl_cache_key not in st.session_state:
        st.session_state[_excl_cache_key] = extract_unique_exclusions(
            _zecom_file, ZECOM_COLS['exclusion'], _excl_start_row
        )
    UNIQUE_EXCLUSIONS = st.session_state[_excl_cache_key]


# ============================================================
# HELPER FUNCTIONS
# ============================================================
def normalise_ean(v):
    return str(v).strip().split('.')[0].strip()


def matches_mechanic(excl_val, mech):
    el = str(excl_val).strip().lower()
    mv = [v.lower() for v in mech['match_values']]
    ex = [v.lower() for v in mech.get('excludes', [])]
    matched = any(m in el for m in mv) if mech['match_type'] == 'contains' else el in mv
    if not matched:
        return False
    if any(x in el for x in ex if x):
        return False
    return True


def build_mechanics_from_ui(raw_mechanics):
    parsed = []
    for m in raw_mechanics:
        name = m['name'].strip()
        mv_raw = m['match_values']
        if isinstance(mv_raw, list):
            match_values = [v.strip() for v in mv_raw if v and v.strip()]
        else:
            match_values = [v.strip() for v in str(mv_raw).splitlines() if v.strip()]
        if not name or not match_values:
            continue
        excludes = [v.strip() for v in m['excludes'].split(',') if v.strip()]
        parsed.append({'name': name, 'match_type': m['match_type'], 'match_values': match_values, 'excludes': excludes})
    return parsed


def run_filtration(price_f, content_f, zecom_f, am_f, mechanics, zecom_cols, data_start_row,
                   apply_threshold, min_price):
    logs = []
    def log(m): logs.append(m)

    price_f.seek(0)
    wb = openpyxl.load_workbook(price_f)
    sheet_name = wb.sheetnames[0]
    ws = wb[sheet_name]
    log(f"▶ Loaded '{sheet_name}' sheet from Price/Stock file ({ws.max_row:,} rows)")

    headers = [c.value for c in ws[1]]
    def find_col(name):
        for idx, h in enumerate(headers, start=1):
            if h and str(h).strip().lower() == name.lower():
                return idx
        return None

    sellersku_col = find_col('SellerSKU')
    price_col = find_col('Price')
    if not sellersku_col:
        raise ValueError("Could not find a 'SellerSKU' column in row 1 of the Price/Stock file.")
    if not price_col:
        raise ValueError("Could not find a 'Price' column in row 1 of the Price/Stock file.")

    last_col = max((idx for idx, h in enumerate(headers, start=1) if h is not None), default=len(headers))

    # ── Content file → EAN_norm -> ALU_NO ───────────────────
    log("▶ Mapping ALU_NO via Content file...")
    content_f.seek(0)
    content_wb = openpyxl.load_workbook(content_f, read_only=True, data_only=True)
    if 'content' not in content_wb.sheetnames:
        raise ValueError("Content file must contain a sheet named 'content'.")
    content_ws = content_wb['content']
    c_headers = [c.value for c in next(content_ws.iter_rows(min_row=1, max_row=1))]
    try:
        ean_i = [str(h).strip().lower() for h in c_headers].index('ean')
        color_i = [str(h).strip().lower() for h in c_headers].index('color_no')
    except ValueError:
        raise ValueError("Content file's 'content' sheet must have 'EAN' and 'Color_No' columns.")
    content_map = {}
    for row in content_ws.iter_rows(min_row=2, values_only=True):
        if row[ean_i] is None:
            continue
        content_map[normalise_ean(row[ean_i])] = str(row[color_i]).strip()
    content_wb.close()
    log(f"   ✓ {len(content_map):,} EAN → ALU_NO mappings loaded")

    # ── zeCOM Tracking file → ALU_NO -> record ──────────────
    log("▶ Loading zeCOM Tracking file (sheet 'MY')...")
    zecom_f.seek(0)
    zecom_wb = openpyxl.load_workbook(zecom_f, read_only=True, data_only=True)
    if 'MY' not in zecom_wb.sheetnames:
        raise ValueError("zeCOM Tracking file must contain a sheet named 'MY'.")
    zecom_ws = zecom_wb['MY']
    zecom_map = {}
    for row in zecom_ws.iter_rows(min_row=5, values_only=True):
        style = row[zecom_cols['style']] if len(row) > zecom_cols['style'] else None
        if not style:
            continue
        def g(key, row=row):
            idx = zecom_cols.get(key)
            return row[idx] if idx is not None and len(row) > idx else None
        zecom_map[str(style).strip()] = {
            'Launch_Date': g('launch_date'),
            'Status_Lazada': g('status_laz'),
            'Price': g('price'),
            'Special_Price': g('special_price'),
            'Disc_Pct': g('disc_pct'),
            'Exclusion': g('exclusion'),
        }
    zecom_wb.close()
    log(f"   ✓ {len(zecom_map):,} styles loaded from zeCOM")

    # ── AM Exclusion file (optional) ────────────────────────
    am_set = set()
    if am_f:
        log("▶ Loading AM Exclusion file...")
        am_f.seek(0)
        am_wb = openpyxl.load_workbook(am_f, read_only=True, data_only=True)
        am_ws = am_wb.active
        for row in am_ws.iter_rows(min_row=2, values_only=True):
            if row and row[0] is not None:
                am_set.add(str(row[0]).strip())
        am_wb.close()
        log(f"   ✓ {len(am_set):,} AM-excluded ALU_NOs")
    else:
        log("▶ AM Exclusion file not provided — no ALU_NO will be AM-excluded")

    # ── Append new headers ───────────────────────────────────
    new_headers = ['ALU_NO', 'Live', 'Status', 'RRP', 'SRP', 'RRP check', '%', 'Exclusions'] \
        + [m['name'] for m in mechanics]
    base_font = Font(name='Aptos Narrow', size=11, bold=False)
    yellow_fill = PatternFill(fill_type='solid', fgColor='FFFFFF00')
    n_lookup_cols = 8
    for i, h in enumerate(new_headers):
        col = last_col + 1 + i
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = base_font
        if i >= n_lookup_cols:
            cell.fill = yellow_fill
        ws.column_dimensions[cell.column_letter].width = max(
            ws.column_dimensions[cell.column_letter].width or 0, 14
        )

    counts = {m['name']: 0 for m in mechanics}
    total = matched_alu = matched_zecom = 0

    for r in range(int(data_start_row), ws.max_row + 1):
        sellersku = ws.cell(row=r, column=sellersku_col).value
        if sellersku is None:
            continue
        total += 1
        orig_price = ws.cell(row=r, column=price_col).value

        alu_no = content_map.get(normalise_ean(sellersku))

        if alu_no is None:
            row_vals = [NA] * n_lookup_cols
        else:
            matched_alu += 1
            rec = zecom_map.get(alu_no)
            if rec is None:
                row_vals = [alu_no] + [NA] * (n_lookup_cols - 1)
            else:
                matched_zecom += 1
                live, status = rec['Launch_Date'], rec['Status_Lazada']
                rrp, srp, pct, excl = rec['Price'], rec['Special_Price'], rec['Disc_Pct'], rec['Exclusion']
                try:
                    rrp_check = round(float(rrp), 2) == round(float(orig_price), 2)
                except (TypeError, ValueError):
                    rrp_check = NA
                row_vals = [alu_no, live, status, rrp, srp, rrp_check, pct, excl]

        status_val, live_val, rrp_val, srp_val, excl_val = row_vals[2], row_vals[1], row_vals[3], row_vals[4], row_vals[7]

        eligible = (status_val == 'YES')
        if isinstance(live_val, (date, datetime)):
            live_d = live_val.date() if isinstance(live_val, datetime) else live_val
            if live_d > TODAY:
                eligible = False
        if alu_no and alu_no in am_set:
            eligible = False
        if apply_threshold:
            for price_val in (rrp_val, srp_val):
                if price_val not in (NA, None):
                    try:
                        if float(price_val) < min_price:
                            eligible = False
                            break
                    except (TypeError, ValueError):
                        pass

        mech_vals = []
        for m in mechanics:
            if excl_val == NA or not eligible:
                mech_vals.append('')
            else:
                yes = matches_mechanic(excl_val, m)
                if yes:
                    counts[m['name']] += 1
                mech_vals.append('YES' if yes else '')

        for i, val in enumerate(row_vals + mech_vals):
            col = last_col + 1 + i
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = base_font
            if i == 1 and isinstance(val, (date, datetime)):
                cell.number_format = 'mm-dd-yy'

    log(f"   ✓ {total:,} SKUs processed | {matched_alu:,} matched to ALU_NO | {matched_zecom:,} matched to zeCOM")
    for name, count in counts.items():
        log(f"   {name}: {count:,} SKUs")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output, logs, counts, total, matched_alu, matched_zecom


# ============================================================
# TAB 1 — RUN FILTRATION
# ============================================================
with tab_run:

    # ── Mechanics ──────────────────────────────────────────
    st.subheader("🎯 Voucher Mechanics")
    st.caption("Mechanics (names, percentages, wording) change every campaign — add, edit, or remove as many as you need.")

    if 'mechanics' not in st.session_state:
        st.session_state.mechanics = [
            {'name': '', 'match_type': 'contains', 'match_values': [], 'excludes': ''}
        ]

    _remove_idx = None
    for i, mech in enumerate(st.session_state.mechanics):
        with st.container(border=True):
            c1, c2, c3 = st.columns([4, 2, 1])
            mech['name'] = c1.text_input(
                "Mechanic name (becomes the output column header)",
                value=mech['name'], key=f"mech_name_{i}",
                placeholder="e.g. 20% NMS (All 20% VC Remark)"
            )
            mech['match_type'] = c2.selectbox(
                "Match type", ['contains', 'exact'],
                index=['contains', 'exact'].index(mech['match_type']),
                key=f"mech_type_{i}"
            )
            c3.markdown("&nbsp;")
            if c3.button("🗑️ Remove", key=f"mech_remove_{i}"):
                _remove_idx = i
            if UNIQUE_EXCLUSIONS:
                _current = mech['match_values'] if isinstance(mech['match_values'], list) else []
                _valid_default = [v for v in _current if v in UNIQUE_EXCLUSIONS]
                mech['match_values'] = st.multiselect(
                    "zeCOM Exclusion value(s) this mechanic matches — fetched from the uploaded zeCOM file",
                    options=UNIQUE_EXCLUSIONS,
                    default=_valid_default,
                    key=f"mech_mv_ms_{i}",
                )
            else:
                _current_text = mech['match_values'] if isinstance(mech['match_values'], str) else "\n".join(mech['match_values'])
                mech['match_values'] = st.text_area(
                    "zeCOM Exclusion value(s) this mechanic matches (one per line)",
                    value=_current_text, key=f"mech_mv_ta_{i}", height=70,
                    placeholder="20% VC",
                    help="Upload the zeCOM Tracking file to pick from actual values instead of typing them."
                )
            mech['excludes'] = st.text_input(
                "Exclude if Exclusion text also contains (comma-separated, optional)",
                value=mech['excludes'], key=f"mech_ex_{i}",
                placeholder="e.g. No platform VC"
            )

    if _remove_idx is not None:
        st.session_state.mechanics.pop(_remove_idx)
        st.rerun()

    if st.button("➕ Add another mechanic"):
        st.session_state.mechanics.append({'name': '', 'match_type': 'contains', 'match_values': [], 'excludes': ''})
        st.rerun()

    st.divider()

    # ── File uploads ───────────────────────────────────────
    st.subheader("📂 Upload Input Files")

    col1, col2, col3 = st.columns(3)
    with col1:
        st.markdown("**1. Price / Stock File** *(required)*")
        price_file = st.file_uploader("SellerPriceTemplate", type=['xlsx', 'xls'], key='price')
    with col2:
        st.markdown("**2. Content File** *(required)*")
        content_file = st.file_uploader("Content_file", type=['xlsx', 'xls'], key='content')
    with col3:
        st.markdown("**3. zeCOM Tracking File** *(required)*")
        zecom_file = st.file_uploader("zeCOM_Tracking_File", type=['xlsx', 'xls'], key='zecom')

    col4, col5 = st.columns(2)
    with col4:
        st.markdown("**4. AM Exclusion File** *(optional)*")
        am_file = st.file_uploader("AM_Exclusion.xlsx", type=['xlsx', 'xls'], key='am')
    with col5:
        st.markdown("**Output filename** *(optional)*")
        output_filename = st.text_input(
            "Leave blank to auto-generate", value="", label_visibility="collapsed",
            placeholder="e.g. Brand_Spotlight_23_June_voucher_filtration.xlsx"
        )

    st.divider()

    # ── Run button ─────────────────────────────────────────
    parsed_mechanics_preview = build_mechanics_from_ui(st.session_state.mechanics)
    ready = price_file and content_file and zecom_file and parsed_mechanics_preview

    if not ready:
        st.info("👆 Upload the 3 required files and define at least one mechanic (with a name and match value) to get started")

    if st.button("▶  Run Voucher Filtration", disabled=not ready):
        with st.spinner("Processing..."):
            try:
                excel_bytes, logs, counts, total, matched_alu, matched_zecom = run_filtration(
                    price_file, content_file, zecom_file, am_file,
                    parsed_mechanics_preview, ZECOM_COLS, data_start_row,
                    apply_threshold, min_price
                )

                st.success("✅ Filtration complete!")
                st.divider()
                st.subheader("📊 Results Summary")

                metrics_cols = st.columns(3 + len(parsed_mechanics_preview))
                with metrics_cols[0]:
                    st.metric("Total SKUs", f"{total:,}")
                with metrics_cols[1]:
                    st.metric("Matched to ALU_NO", f"{matched_alu:,}")
                with metrics_cols[2]:
                    st.metric("Matched to zeCOM", f"{matched_zecom:,}")
                for i, (name, count) in enumerate(counts.items()):
                    with metrics_cols[3 + i]:
                        st.metric(name[:24], f"{count:,}")

                with st.expander("📝 Processing Log"):
                    for line in logs:
                        st.text(line)

                st.divider()
                default_name = output_filename.strip() or (
                    price_file.name.rsplit('.', 1)[0] + "_voucher_filtration.xlsx"
                )
                st.download_button(
                    label="⬇️  Download Filtered Excel",
                    data=excel_bytes,
                    file_name=default_name,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )

            except Exception as e:
                st.error(f"❌ Error: {e}")
                st.exception(e)


# ============================================================
# TAB 2 — CLEAR OUTPUT COLUMNS
# ============================================================
with tab_clear:

    st.subheader("🗑️ Clear Output Columns from Processed File")
    st.caption(
        "Strips the appended columns (ALU_NO, Live, Status, RRP, SRP, RRP check, %, "
        "Exclusions, and any mechanic YES/blank columns) from a file that was already "
        "processed — returning it to the original clean Lazada SellerPriceTemplate."
    )
    st.divider()

    clear_file = st.file_uploader(
        "Upload the processed Excel file to clear",
        type=['xlsx', 'xls'],
        key='clear_upload'
    )

    if clear_file:
        clear_file.seek(0)
        try:
            wb_preview = openpyxl.load_workbook(clear_file, read_only=True)
        except Exception as e:
            st.error(f"Could not open file: {e}")
            wb_preview = None

        if wb_preview:
            ws_preview = wb_preview[wb_preview.sheetnames[0]]
            row1 = [ws_preview.cell(row=1, column=c).value for c in range(1, ws_preview.max_column + 1)]
            wb_preview.close()

            strip_from = None
            for idx, h in enumerate(row1, start=1):
                if h and str(h).strip().upper() == 'ALU_NO':
                    strip_from = idx
                    break

            if strip_from:
                total_cols = len(row1)
                n_to_strip = total_cols - strip_from + 1
                stripped_names = [str(row1[c - 1]) for c in range(strip_from, total_cols + 1)]

                st.success(
                    f"✅ Found **ALU_NO** at column **{get_column_letter(strip_from)}** (col {strip_from}). "
                    f"Ready to remove **{n_to_strip}** column(s)."
                )

                with st.expander("Columns that will be removed"):
                    for name in stripped_names:
                        st.markdown(f"- `{name}`")

                st.divider()

                if st.button("🗑️  Strip columns and prepare download", use_container_width=True, key='do_clear'):
                    with st.spinner("Removing output columns..."):
                        clear_file.seek(0)
                        wb_out = openpyxl.load_workbook(clear_file)
                        ws_out = wb_out[wb_out.sheetnames[0]]
                        ws_out.delete_cols(strip_from, n_to_strip)
                        buf = io.BytesIO()
                        wb_out.save(buf)
                        buf.seek(0)

                    clean_name = (
                        clear_file.name
                        .replace('_voucher_filtration', '')
                        .rsplit('.', 1)[0] + '_clean.xlsx'
                    )
                    st.download_button(
                        label="⬇️  Download Clean File",
                        data=buf,
                        file_name=clean_name,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                        key='download_clean'
                    )
                    st.success(
                        f"✅ Done! Removed **{n_to_strip}** column(s) — file is ready to download."
                    )

            else:
                st.warning(
                    "⚠️ No **ALU_NO** column found in row 1 of this file. "
                    "It may already be clean, or the header row is in an unexpected position."
                )
